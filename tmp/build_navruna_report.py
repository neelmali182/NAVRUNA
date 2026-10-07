from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = r"C:\Users\neelm\Downloads\PROJECJ CLG\NAVRUNA\NAVRUNA_Project_Report.docx"
NAVY = "17365D"
MIDBLUE = "DCE6F1"
PALE = "F3F6FA"
GRAY = "D9D9D9"
BLACK = "000000"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.78)
sec.bottom_margin = Inches(.72)
sec.left_margin = Inches(.82)
sec.right_margin = Inches(.82)
sec.header_distance = Inches(.32)
sec.footer_distance = Inches(.32)
sec.different_first_page_header_footer = True

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12

for name, size in [("Title", 25), ("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11.5)]:
    st = styles[name]
    st.font.name = "Times New Roman"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    st.font.size = Pt(size)
    st.font.bold = name != "Title"
    st.font.color.rgb = RGBColor.from_string(BLACK)
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
    st.paragraph_format.space_after = Pt(5)
styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
styles["Title"].paragraph_format.space_after = Pt(10)

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=95, start=110, bottom=95, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "5")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), GRAY)

def add_table(headers, rows, widths=None, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    if widths:
        for c, width in zip(table.columns, widths):
            c.width = Inches(width)
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        if widths: cell.width = Inches(widths[i])
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell, 110, 115, 110, 115)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        r.bold = True; r.font.name = "Times New Roman"; r.font.size = Pt(font_size); r.font.color.rgb = RGBColor(255,255,255)
    for idx, row in enumerate(rows):
        cells = table.add_row().cells
        for j, text in enumerate(row):
            if widths: cells[j].width = Inches(widths[j])
            set_cell_margins(cells[j])
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if idx % 2 == 1: set_cell_shading(cells[j], PALE)
            p = cells[j].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(str(text)); r.font.name = "Times New Roman"; r.font.size = Pt(font_size)
    return table

PENDING_PAGE_BREAK = False
def h1(text):
    global PENDING_PAGE_BREAK
    p = doc.add_heading(text, level=1)
    if PENDING_PAGE_BREAK:
        p.paragraph_format.page_break_before = True
        PENDING_PAGE_BREAK = False
    return p
def h2(text): doc.add_heading(text, level=2)
def para(text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = True
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else: p.add_run(text)
    return p
def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p
def numbered(text, number):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(.28)
    p.paragraph_format.first_line_indent = Inches(-.28)
    p.paragraph_format.space_after = Pt(3)
    p.add_run(f"{number}.\t{text}")
    return p
def page():
    global PENDING_PAGE_BREAK
    PENDING_PAGE_BREAK = True

# Running header and page number are omitted from the cover.
header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = header.add_run("NAVRUNA  |  OFFLINE MARITIME AI SIMULATION LAB")
hr.font.name = "Times New Roman"; hr.font.size = Pt(8); hr.font.color.rgb = RGBColor.from_string("555555")
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run("NAVRUNA PROJECT REPORT  •  ")
fr.font.name = "Times New Roman"; fr.font.size = Pt(8); fr.font.color.rgb = RGBColor.from_string("555555")
field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

# Cover page
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(68); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("NAVRUNA"); r.bold = True; r.font.name = "Times New Roman"; r.font.size = Pt(35); r.font.color.rgb = RGBColor.from_string(NAVY)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Offline Maritime AI Simulation Lab"); r.font.name = "Times New Roman"; r.font.size = Pt(18); r.font.color.rgb = RGBColor.from_string(BLACK)
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(15); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("A Comprehensive Technical Report"); r.italic = True; r.font.name = "Times New Roman"; r.font.size = Pt(14)
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(70); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("PROJECT BASED LEARNING REPORT"); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = RGBColor.from_string(NAVY)
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(22); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Prepared by\nNAVRUNA Project Team").font.size = Pt(12)
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(42); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("October 2026"); r.font.size = Pt(11); r.font.color.rgb = RGBColor.from_string("555555")
page()

h1("Abstract")
para("Maritime route planning must account for destination progress, fuel limits, changing ocean conditions, and the risk of land collision. NAVRUNA is an offline-first simulation lab for studying these challenges with reinforcement learning. Unlike a route-following demonstration, the project gives each simulated vessel an origin, a destination, a fuel budget, and a local synthetic ocean state. A shared Proximal Policy Optimization (PPO) policy selects heading changes from observations at each simulation step.")
para("The software combines a FastAPI backend, a PyTorch PPO trainer, a multi-vessel simulator, a bundled port catalog and coastline dataset, and a browser interface built around CesiumJS. Each policy observation contains 25 normalized values describing vessel state, destination geometry, fuel, weather, and nearby land clearance. The policy chooses among seven discrete heading changes. A separate geometry-based safety layer screens movement and checks the swept path against coastline data, so the learned policy does not have final authority over land collision constraints. Procedural weather fields provide spatially and temporally coherent wind, waves, currents, storms, precipitation, visibility, and pressure without requiring a live data feed.")
para("The repository documents an implemented research prototype with training telemetry, fleet analytics, a weather visualizer, and local simulation assets. It does not provide a completed benchmark study establishing navigation success, fuel savings, or real-world safety. Accordingly, this report describes the implemented design and its evaluation approach without presenting unmeasured performance claims. NAVRUNA is intended for software engineering education and simulation research, not for controlling real vessels.")
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(5)
r=p.add_run("Keywords: "); r.bold=True
p.add_run("maritime simulation, reinforcement learning, PPO, autonomous navigation, collision avoidance, procedural weather")

h1("Report Organization")
para("The report introduces the navigation problem and project objectives, reviews relevant reinforcement-learning and geospatial concepts, and then explains NAVRUNA's architecture and methods. Implementation details are followed by an evidence-based analysis of what the repository supports, the limits of current evaluation, and directions for further work.")
page()

h1("1 Introduction")
h2("1.1 Background")
para("Vessel routing is a sequential decision problem. A vessel must make progress toward a destination while its heading, fuel reserve, nearby coastline, and environmental conditions change over time. A static line between two ports does not represent the repeated decisions needed to navigate under local constraints. Reinforcement learning offers a way to train a policy by interacting with a simulated environment and receiving feedback from its actions.")
para("NAVRUNA turns this idea into an interactive, multi-vessel learning environment. It is designed to run locally with packaged port and land data and deterministic synthetic weather. Users can launch a fleet, advance simulation time, observe telemetry on a globe, and run PPO training while viewing learning diagnostics.")
h2("1.2 Problem Statement")
para("Educational demonstrations of autonomous navigation often focus on a learned policy without making environmental assumptions, safety constraints, and model-training behavior visible. Other route visualizations show a precomputed path, which can obscure the decision process. NAVRUNA addresses this gap by providing an inspectable simulation in which a policy selects incremental heading changes, environmental conditions are explicit, and a separate safety shield validates movement against coastline geometry.")
h2("1.3 Significance")
para("The project brings together reinforcement learning, geospatial data processing, procedural environmental modeling, API design, and interactive visualization in one software system. Its main contribution is a reproducible classroom and research platform for examining how a navigation policy behaves under synthetic conditions and hard geometric constraints. The distinction between a learned proposal and an authoritative safety check also makes an important systems-engineering principle visible: safety-critical constraints should not depend solely on a policy's learned behavior.")
h2("1.4 Report Scope")
para("This report covers the current software visible in the NAVRUNA repository and its accompanying project documentation. It explains the simulator, PPO training path, coastline checks, backend endpoints, and browser interface. It does not claim that the synthetic environment is a validated model of real ocean conditions or that the simulator's safety checks are suitable for operational navigation.")
page()

h1("2 Objectives and Scope")
h2("2.1 Project Objectives")
for n,item in enumerate([
    "Create an offline-first global maritime simulation using a local port catalog and bundled coastline data.",
    "Represent vessel navigation as sequential heading decisions instead of supplying the policy with a precomputed route polyline.",
    "Train a shared PPO policy from normalized vessel, destination, fuel, weather, and land-clearance observations.",
    "Enforce coastline constraints through a deployment-time action shield that evaluates candidate swept movements.",
    "Model wind, waves, currents, storms, precipitation, visibility, and pressure with a deterministic synthetic weather field.",
    "Expose simulation, training, and weather information through a FastAPI service and present it in a browser-based globe and training dashboard.",
    "Record useful simulation and training telemetry so behavior can be inspected and future experiments can be compared."
], 1): numbered(item, n)
h2("2.2 System Scope")
add_table(["Included in the prototype", "Outside the validated scope"], [
    ("Synthetic port-to-port voyages and multi-vessel fleet simulation", "Live vessel control or operational navigation advice"),
    ("Bundled port and coastline data; local weather generation", "Live AIS, forecast, bathymetry, or nautical-chart integration"),
    ("PPO training, telemetry, checkpoints, and policy inference", "A published benchmark proving generalization or fuel savings"),
    ("Action masking and a 5 km coastline safety buffer in simulation", "COLREGS compliance, certified collision avoidance, or maritime assurance"),
    ("CesiumJS visualization, scenario controls, and analytics", "A fully air-gapped map stack; the globe references CDN assets and an online basemap")
], widths=[3.2,3.2], font_size=9.4)
h2("2.3 Intended Users")
para("The intended users are students, developers, and researchers exploring reinforcement learning and simulation software. The project supports controlled experimentation with synthetic scenarios and inspection of training behavior. It is not intended to replace a qualified navigator, charting system, or vessel safety system.")
page()

h1("3 Review of Related Concepts")
h2("3.1 Reinforcement Learning for Sequential Control")
para("In reinforcement learning, an agent observes a state, chooses an action, and receives a reward and a new state from an environment. This interaction is repeated over an episode. For navigation, the agent's state can describe vessel position and heading, destination direction and distance, remaining fuel, weather, and local obstacles. The reward can encourage progress while penalizing resource use and unsafe outcomes. The quality of a learned policy depends on the environment, reward definition, observation design, and range of training scenarios.")
h2("3.2 Proximal Policy Optimization")
para("Proximal Policy Optimization is a policy-gradient method that updates a stochastic policy using a clipped objective intended to limit overly large changes between successive policies. The method is commonly paired with a value function and generalized advantage estimation. NAVRUNA implements PPO with a shared neural-network body and separate policy and value outputs. Its trainer collects rollouts from parallel environments, masks actions during training, computes advantages, and performs several optimization epochs per update [1, 2].")
h2("3.3 Geospatial Data and Collision Constraints")
para("A geographic safety check must account for the path traveled between two positions, not only the endpoints. NAVRUNA packages coastline polygons and uses both a raster distance field for fast screening and vector geometry for final validation. The global dataset is derived from GSHHG, a hierarchical shoreline and geographic database [3]. This hybrid method supports fast policy observations while retaining a more detailed geometric check along proposed movement segments.")
h2("3.4 Interactive Geospatial Visualization")
para("CesiumJS provides a browser library for interactive 3D globes and maps [4]. In NAVRUNA, the globe displays vessel positions and status; separate panels expose fleet measures, vessel telemetry, weather layers, and training controls. This presentation connects policy decisions to their spatial context and makes the simulation easier to inspect than a terminal-only workflow.")
h2("3.5 Design Implications")
para("Three design choices follow from these concepts. First, the policy receives local state rather than a prescribed route. Second, the simulated world provides repeatable, synthetic inputs for controlled learning experiments. Third, collision screening remains a distinct software layer, so the learned policy may propose an unsafe action without that action being executed in the deployment simulation. These choices improve experiment visibility but do not, by themselves, demonstrate real-world validity.")
page()

h1("4 System Architecture")
para("NAVRUNA is organized as a local web application with a simulation API, simulation and training models, bundled geospatial data, and a browser interface. The backend initializes computationally heavy components lazily so the health endpoint can respond before the coastline grid and PyTorch model are loaded.")
add_table(["Layer", "Implementation", "Responsibility"], [
    ("User interface", "HTML, CSS, JavaScript, CesiumJS", "Fleet controls, globe, telemetry, weather layers, analytics, training charts"),
    ("HTTP API", "FastAPI", "Simulation lifecycle, fleet state, analytics, ports, weather, policy and training routes"),
    ("Fleet simulator", "Python", "Creates vessels, advances voyages, applies decisions, updates metrics and events"),
    ("Navigation policy", "PyTorch PPO", "Maps 25-value observations to seven heading-change scores; estimates state value"),
    ("Safety and geography", "Shapely, NumPy, SciPy", "Loads land polygons, builds distance grid, screens actions, verifies swept segments"),
    ("Weather model", "Deterministic procedural Python model", "Produces coherent ocean conditions at a location and simulation time"),
    ("Local data", "JSON and GeoJSON assets", "Port locations, coastline polygons and frontend copies of simulation data")
], widths=[1.15,1.8,3.45], font_size=8.8)
h2("4.1 Request and Simulation Flow")
for n,item in enumerate([
    "The browser sends a request to start a scenario or advance the simulator through the /api/simulation routes.",
    "The fleet simulator obtains each active vessel's observation and asks the policy for candidate heading scores. Before movement, it validates actions against the coastline safety geometry.",
    "The environment updates vessel position, heading, fuel, progress, weather values, and episode outcomes. Completed voyages are assigned a new destination; failures and collision events are recorded.",
    "The API returns fleet state and analytics to the browser, which refreshes the globe and telemetry panels. The training page separately polls training progress and history."
], 1): numbered(item, n)
h2("4.2 Principal API Routes")
add_table(["Route", "Purpose"], [
    ("GET /api/health", "Liveness and offline-mode status"),
    ("POST /api/simulation/start", "Generate a fleet with a selected vessel count and optional seed"),
    ("POST /api/simulation/step", "Advance simulated time"),
    ("GET /api/simulation/state", "Return the current vessel snapshot and metrics"),
    ("GET /api/simulation/analytics", "Return fleet and trainer analytics"),
    ("GET /api/simulation/weather", "Return a procedural global weather field"),
    ("POST /api/simulation/train", "Start PPO training in a background worker"),
    ("GET /api/simulation/training", "Return current training status, report, history, and logs")
], widths=[2.15,4.25], font_size=9)
page()

h1("5 Research Methodology")
h2("5.1 Research Design")
para("The project follows an experimental software-engineering design. The environment, learning algorithm, safety system, API, and visualization are separate enough to inspect individually while remaining integrated in one simulation workflow. A scenario seed can be supplied when creating a fleet, which supports repeatable simulation starts. Training environments use configured seeds and a staged curriculum.")
h2("5.2 Environment State and Action Space")
para("Each vessel is modeled as an episode with an origin and destination selected from the local port catalog. The environment positions each endpoint offshore and initializes heading, speed, and fuel. The policy receives a 25-dimensional normalized observation. It includes position and heading, relative destination bearing and distance, fuel fraction, speed, progress, local wind and current direction, wave height, storm intensity, seven directional land-clearance probes, and normalized step count.")
add_table(["Element", "Project configuration"], [
    ("Observation", "25 normalized values for vessel, target, fuel, environment, and land proximity"),
    ("Action", "Seven discrete heading changes: -35°, -20°, -10°, 0°, +10°, +20°, +35°"),
    ("Episode success", "Destination reached within 20 nautical miles"),
    ("Terminal failure", "Land collision, fuel exhaustion, or configured time/step limit"),
    ("Terminal rewards", "+2 for destination reached; -2 for collision or fuel/time failure"),
    ("Intermediate shaping", "Small reward for distance progress, with fuel-use and low-clearance penalties")
], widths=[1.55,4.85], font_size=9)
h2("5.3 Training Procedure")
para("The PPO trainer creates parallel environments, collects rollout observations and rewards, and uses a shared policy/value network with 256-, 256-, and 128-unit hidden layers. The policy has seven output logits and the value head has one output. Training uses a learning rate default of 0.0003, discount factor 0.99, GAE lambda 0.95, a clipping range of 0.2, entropy coefficient of 0.01, and gradient clipping. The curriculum begins with shorter routes and gradually increases route length; land constraints are enabled after the early open-water phase. Training status and update metrics are written to local history and checkpoint files.")
h2("5.4 Safety Method")
para("The policy's action is checked before execution. A fast raster-based mask screens likely unsafe directions, then vector coastline geometry checks the swept movement segment against a configurable 5 km buffer. The fleet simulator uses the verified candidates and does not execute a movement when no safe candidate remains. During training, collision can also terminate an episode with a penalty; masking is used to reduce unsafe samples. The safety guarantee applies only to this simulation and the resolution and accuracy of its bundled data.")
h2("5.5 Weather and Data")
para("The weather simulator deterministically calculates conditions from latitude, longitude, time, and a seed. It supplies wind speed and direction, wave height, current speed and direction, storm intensity, precipitation, visibility, pressure, and sea state. The bundled port catalog contains 279 locations. Coastline polygons are stored in a local GeoJSON asset, and the backend builds a 0.25-degree land and distance grid for fast queries. The frontend also contains local copies of port and land assets.")

h1("6 Implementation Details")
h2("6.1 Backend and Simulation API")
para("The FastAPI application serves API routes and the static frontend. The health endpoint reports the offline synthetic mode without importing heavy model dependencies. Simulation, weather, and trainer objects are initialized on demand. Fleet stepping is moved to worker threads so synchronous simulation work does not block the async request loop. Training is submitted to a single-worker executor, and a lock reserves the trainer's starting state to prevent duplicate concurrent submissions.")
h2("6.2 Geospatial Safety")
para("The geography module loads polygon features, builds their union, and prepares the geometry for repeated spatial queries. A raster grid supports quick land lookup and approximate distance-to-land queries. For movement validation, the simulator samples points along the proposed segment and checks both direct polygon intersection and proximity to the safety buffer. This continuous segment check is designed to prevent a vessel from skipping across a narrow land feature between safe endpoints.")
h2("6.3 Fleet and Voyage Management")
para("A scenario can contain between 1 and 300 vessels through the API request model; the interface defaults to 60. A simulation step advances the fleet in hourly increments and supports fractional time steps. Each vessel maintains navigation state, fuel use, weather telemetry, risk indicators, progress, and its most recent action. A successful voyage increments completion metrics and assigns a new port pair. A collision, no-safe-action result, fuel exhaustion, or time limit is recorded as a failure event.")
h2("6.4 Browser Interface")
para("The main screen combines a 3D globe with mission controls, a selectable vessel inspector, fleet outcomes, and environment telemetry. The weather visualizer can show waves, wind, currents, or storms. A separate Training Lab presents learning status, episode outcomes, training speed, losses, entropy, approximate KL, learning-rate information, and activity logs. The UI also reports model device information and configuration.")
h2("6.5 Running the Prototype")
para("On Windows, the project provides run_navruna.bat and start.ps1 launch paths. The launcher starts the API and frontend locally and opens the browser interface. The project README documents the usual local URLs and dependencies. The user can start a simulation, alter fleet size and simulated time, inspect vessel state, open weather layers, and begin or stop training through the interface.")
page()

h1("7 Data Analysis and Interpretation")
h2("7.1 What Can Be Established from the Repository")
para("The implementation establishes the following design properties: vessel actions are selected step by step; the policy receives a fixed-size observation and discrete action set; the trainer supports parallel PPO rollouts and curriculum progression; the simulation checks candidate paths against coastline geometry; and the browser exposes fleet, weather, and training telemetry. The automated test source includes checks for the port catalog, land geometry, observation and network dimensions, no-preplanned-route fleet behavior, fractional time stepping, voyage reassignment, and no-safe-action failure handling.")
h2("7.2 Evaluation Measures")
para("A full empirical evaluation should record voyage success rate, collision count, no-safe-action count, fuel use, distance traveled, route completion time, and policy reward across repeated seeds. Training evaluation should use held-out origin-destination pairs and report uncertainty across multiple training runs. Baselines should include direct-to-destination steering with the same safety shield and a classical path-planning method. Stress scenarios should vary route length, coastline complexity, weather severity, and fuel budget.")
add_table(["Measure", "What it would show", "Evidence needed"], [
    ("Voyage success rate", "Ability to reach destinations under defined conditions", "Completed and failed episodes over repeated held-out scenarios"),
    ("Collision and safety events", "Whether the shield prevents unsafe executed movement", "Collision counts and records of rejected actions by scenario"),
    ("Fuel and time use", "Resource efficiency compared with a baseline", "Consistent fuel model, elapsed simulated time, and baseline runs"),
    ("Training stability", "Learning progress and sensitivity to random seeds", "Learning curves, checkpoint metadata, and repeated runs"),
    ("Runtime performance", "Fleet and API responsiveness as scale increases", "Timed runs with documented hardware, fleet size, and step size")
], widths=[1.45,2.2,2.75], font_size=8.8)
h2("7.3 Interpretation and Evidence Limits")
para("The current repository provides an implemented simulator and a test suite definition, but no supplied, reproducible benchmark dataset or report of completed PPO evaluation runs. The training dashboard is an observability feature, not a performance result. Therefore, this report does not assign a success percentage, claim fuel savings, or compare NAVRUNA numerically with another navigator. Results should be added only after a controlled experiment records configuration, hardware, seeds, training checkpoint, evaluation cases, and raw outcomes.")
page()

h1("8 Limitations and Future Work")
h2("8.1 Current Limitations")
for item in [
    "Weather is procedural and synthetic. It is suitable for repeatable experiments but does not reproduce forecast uncertainty or observed ocean dynamics.",
    "The bundled coastline data and raster safety field have finite resolution. Geographic error, projection effects, and omitted bathymetry can affect apparent clearance.",
    "The simulator uses simplified vessel speed, fuel, and environmental-drift models. These do not capture vessel-specific propulsion, loading, maneuverability, or sea-state response.",
    "The project does not implement COLREGS-aware interactions, traffic separation schemes, port approach procedures, multi-vessel collision avoidance, or human-supervised operational workflows.",
    "The main UI currently references external CesiumJS assets and an online basemap, so simulation data are local but the complete visual stack is not fully air-gapped.",
    "No supplied results establish that a trained policy generalizes to unseen routes, weather conditions, or geographic regions."
]: bullet(item)
h2("8.2 Future Work")
for n,item in enumerate([
    "Build a repeatable evaluation suite with held-out routes, multiple random seeds, explicit baseline policies, and confidence intervals.",
    "Add bathymetry, vessel draft, lane constraints, port arrival windows, and vessel-specific dynamics to the environment.",
    "Replace or supplement procedural weather with replayable historical datasets while preserving deterministic experiment modes.",
    "Model COLREGS rules, traffic interactions, and uncertainty in sensor and geographic information before studying operationally relevant scenarios.",
    "Track experiments and checkpoints with configuration hashes so results can be reproduced and compared.",
    "Self-host the CesiumJS distribution and imagery assets for installations that need a fully air-gapped interface."
], 1): numbered(item, n)
h2("8.3 Research Boundary")
para("NAVRUNA is a synthetic simulation and software-engineering project. Its coastline shield is a prototype software constraint, not a certified navigation safety system. Real-world use would require validated charts and bathymetry, vessel dynamics, COLREGS-aware behavior, uncertainty modeling, independent verification and validation, and qualified human oversight.")
page()

h1("9 Conclusion")
para("NAVRUNA brings together a local maritime simulation, a PPO learning system, a geometry-based coastline shield, procedural weather, a FastAPI backend, and a browser-based geospatial interface. Its stepwise action model makes the policy's navigation task explicit, while training telemetry and vessel inspection tools help users understand the simulation's behavior. The architecture is suitable for classroom demonstrations and controlled research into reinforcement-learning navigation.")
para("The current evidence supports a description of what the software implements, not a claim that the learned policy is already effective or operationally safe. A rigorous next step is to evaluate repeatable held-out scenarios against transparent baselines and report navigation, safety, fuel, time, and runtime metrics. Keeping those conclusions tied to measured results will strengthen NAVRUNA as a research prototype and provide a clear path for future improvements.")

h1("References")
refs = [
    "[1] Schulman, J., Wolski, F., Dhariwal, P., Radford, A., and Klimov, O. (2017). Proximal Policy Optimization Algorithms. arXiv:1707.06347. https://arxiv.org/abs/1707.06347",
    "[2] Schulman, J., Moritz, P., Levine, S., Jordan, M., and Abbeel, P. (2016). High-Dimensional Continuous Control Using Generalized Advantage Estimation. arXiv:1506.02438. https://arxiv.org/abs/1506.02438",
    "[3] Wessel, P. and Smith, W. H. F. (1996). A Global, Self-consistent, Hierarchical, High-resolution Shoreline Database. Journal of Geophysical Research: Solid Earth, 101(B4), 8741-8743. https://doi.org/10.1029/96JB00104",
    "[4] Cesium. CesiumJS Fundamentals. Official learning documentation. https://cesium.com/learn/cesiumjs-fundamentals/",
    "[5] FastAPI. FastAPI Documentation. https://fastapi.tiangolo.com/",
    "[6] NAVRUNA project repository. README, architecture notes, PPO model card, source code, and automated test definitions. Project files supplied with this report, accessed October 2026."
]
for ref in refs:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(.25)
    p.paragraph_format.first_line_indent = Inches(-.25)
    p.paragraph_format.space_after = Pt(7)
    p.add_run(ref)

doc.core_properties.title = "NAVRUNA Offline Maritime AI Simulation Lab"
doc.core_properties.subject = "Project Based Learning Technical Report"
doc.core_properties.author = "NAVRUNA Project Team"
doc.core_properties.keywords = "NAVRUNA, maritime simulation, reinforcement learning, PPO"
doc.save(OUT)
print(OUT)
